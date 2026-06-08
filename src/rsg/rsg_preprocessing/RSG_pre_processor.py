"""Single-file ROS 2 preprocessing node for the Risk Scene Graph pipeline.

The module contains the full implementation of the ``RSG_pre_processor`` node.
It keeps all helper classes in one file to simplify integration while still
separating small responsibilities into pdoc-friendly classes.

The node receives RGB, aligned depth, camera calibration, and odometry streams
from either a replayed ROS 2 bag or live ROS 2 topics. It prepares one
synchronized RGB-D-pose frame and publishes it as ``rsg/msg/RsgFrame`` for the
main worker node.
"""

from __future__ import annotations

import copy
import json
import math
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import cv2
import message_filters
import numpy as np
import rclpy
import yaml
from cv_bridge import CvBridge
from geometry_msgs.msg import PoseStamped
from nav_msgs.msg import Odometry
from rclpy.node import Node
from rclpy.qos import QoSProfile, ReliabilityPolicy, HistoryPolicy, qos_profile_sensor_data
from sensor_msgs.msg import CameraInfo, Image, Imu
from std_msgs.msg import Float64MultiArray, String

from rsg.msg import RsgFrame


def stamp_to_float(stamp: Any) -> float:
    """Convert a ROS time stamp message to seconds as a floating-point value.

    Args:
        stamp: ROS message stamp with ``sec`` and ``nanosec`` fields.

    Returns:
        Stamp time in seconds.
    """
    return float(stamp.sec) + float(stamp.nanosec) * 1e-9


@dataclass
class PreprocessorConfig:
    """Configuration values required by the preprocessing node.

    The configuration is loaded from the central YAML pipeline file. Only the
    ``preprocessing`` section is interpreted by this node.
    """

    source: str
    use_sim_time: bool

    rgb_topic: str
    depth_topic: str
    camera_info_topic: str
    odom_topic: str
    imu_topic: str
    frame_topic: str
    status_topic: str

    rgb_depth_slop_sec: float
    odom_buffer_size: int
    odom_tolerance_sec: float
    use_odom_interpolation: bool
    reject_unsynchronized_frames: bool

    use_camera_imu: bool
    imu_buffer_size: int
    imu_tolerance_sec: float
    require_imu_for_frame: bool

    output_rgb_encoding: str
    output_depth_encoding: str
    depth_scale_to_meter: float
    min_depth_m: float
    max_depth_m: float
    max_invalid_depth_ratio: float
    reject_resolution_mismatch: bool

    world_frame: str
    base_frame: str
    camera_frame: str
    base_to_camera_translation: Tuple[float, float, float]
    base_to_camera_rpy: Tuple[float, float, float]

    frame_id_prefix: str
    session_date: str

    sensor_qos_depth: int
    odom_qos_depth: int
    output_qos_depth: int

    global_debug_enabled: bool
    node_debug_enabled: bool
    timing_measurement_enabled: bool
    publish_timing_topic: bool
    timing_topic: str
    write_timing_excel: bool
    timing_excel_path: str
    timing_sheet_name: str
    timing_excel_autosave_every: int

    @property
    def debug_enabled(self) -> bool:
        """Return whether global or node-specific debug mode is enabled."""
        return self.global_debug_enabled or self.node_debug_enabled

    @property
    def timing_enabled(self) -> bool:
        """Return whether preprocessing timing measurement should run."""
        return self.debug_enabled and self.timing_measurement_enabled

    @staticmethod
    def from_yaml(path: str) -> "PreprocessorConfig":
        """Load preprocessing configuration from a YAML file.

        Args:
            path: Absolute or relative path to the central pipeline YAML file.

        Returns:
            Parsed preprocessing configuration.

        Raises:
            FileNotFoundError: If the configuration file does not exist.
            KeyError: If the ``preprocessing`` section is missing.
        """
        config_path = Path(path).expanduser().resolve()
        if not config_path.exists():
            raise FileNotFoundError(f"Configuration file not found: {config_path}")

        with config_path.open("r", encoding="utf-8") as stream:
            root = yaml.safe_load(stream)

        preprocessing = root["preprocessing"]
        # Debug switches may be kept either at the YAML root or inside the
        # preprocessing section. Root-level values are preferred because the
        # file is intended to become the central pipeline configuration.
        root_debug = root.get("debug", {}) or {}
        preprocessing_debug = preprocessing.get("debug", {}) or {}
        debug = {**preprocessing_debug, **root_debug}
        runtime = preprocessing.get("runtime", {})
        topics = preprocessing.get("topics", {})
        sync = preprocessing.get("synchronization", {})
        imu = preprocessing.get("imu", {})
        image = preprocessing.get("image", {})
        frames = preprocessing.get("frames", {})
        frame_id = preprocessing.get("frame_id", {})
        qos = preprocessing.get("qos", {})
        performance = preprocessing.get("performance", {})

        base_to_camera = frames.get("base_to_camera", {})
        translation = tuple(float(v) for v in base_to_camera.get("translation_m", [0.0, 0.0, 0.0]))
        rotation_rpy = tuple(float(v) for v in base_to_camera.get("rotation_rpy_rad", [0.0, 0.0, 0.0]))

        session_date = str(frame_id.get("session_date", ""))
        if not session_date:
            session_date = datetime.now().strftime("%Y%m%d")

        global_debug_enabled = bool(
            debug.get("global_debug", False)
            or debug.get("enabled", False)
        )
        node_debug_enabled = bool(
            debug.get("RSG_pre_processor", False)
            or debug.get("rsg_pre_processor", False)
            or debug.get("preprocessor", False)
        )

        # Each node/run writes a separate Excel workbook to avoid one large
        # shared debug file. The creation timestamp is generated once when the
        # node starts and can be used in the configured filename.
        creation_time = datetime.now().strftime("%H%M%S")
        timing_excel_path = str(
            performance.get(
                "timing_excel_path",
                "~/rsg_ros2_ws/debug/RSG_pre_processor_debug_{session_date}_{creation_time}.xlsx",
            )
        ).format(session_date=session_date, creation_time=creation_time)

        timing_sheet_name = str(performance.get("timing_sheet_name", "timing_debug"))

        return PreprocessorConfig(
            source=str(runtime.get("source", "rosbag")),
            use_sim_time=bool(runtime.get("use_sim_time", True)),
            rgb_topic=str(topics.get("rgb", "/go1/d455/color/image_raw")),
            depth_topic=str(topics.get("depth", "/go1/d455/aligned_depth_to_color/image_raw")),
            camera_info_topic=str(topics.get("camera_info", "/go1/d455/color/camera_info")),
            odom_topic=str(topics.get("odom", "/odom")),
            imu_topic=str(topics.get("imu", imu.get("topic", "/go1/d455/imu"))),
            frame_topic=str(topics.get("prepared_frame", "/rsg/preprocessed/frame")),
            status_topic=str(topics.get("status", "/rsg/preprocessor/status")),
            rgb_depth_slop_sec=float(sync.get("rgb_depth_slop_sec", 0.03)),
            odom_buffer_size=int(sync.get("odom_buffer_size", 300)),
            odom_tolerance_sec=float(sync.get("odom_tolerance_sec", 0.05)),
            use_odom_interpolation=bool(sync.get("use_odom_interpolation", True)),
            reject_unsynchronized_frames=bool(sync.get("reject_unsynchronized_frames", True)),
            use_camera_imu=bool(imu.get("enabled", False)),
            imu_buffer_size=int(imu.get("buffer_size", 1000)),
            imu_tolerance_sec=float(imu.get("tolerance_sec", 0.03)),
            require_imu_for_frame=bool(imu.get("require_for_frame", False)),
            output_rgb_encoding=str(image.get("output_rgb_encoding", "rgb8")),
            output_depth_encoding=str(image.get("output_depth_encoding", "32FC1")),
            depth_scale_to_meter=float(image.get("depth_scale_to_meter", 0.001)),
            min_depth_m=float(image.get("min_depth_m", 0.2)),
            max_depth_m=float(image.get("max_depth_m", 6.0)),
            max_invalid_depth_ratio=float(image.get("max_invalid_depth_ratio", 0.70)),
            reject_resolution_mismatch=bool(image.get("reject_resolution_mismatch", True)),
            world_frame=str(frames.get("world_frame", "odom")),
            base_frame=str(frames.get("base_frame", "base_link")),
            camera_frame=str(frames.get("camera_frame", "d455_color_optical_frame")),
            base_to_camera_translation=translation,
            base_to_camera_rpy=rotation_rpy,
            frame_id_prefix=str(frame_id.get("prefix", "rsg")),
            session_date=session_date,
            sensor_qos_depth=int(qos.get("sensor_depth", 10)),
            odom_qos_depth=int(qos.get("odom_depth", 200)),
            output_qos_depth=int(qos.get("output_depth", 10)),
            global_debug_enabled=global_debug_enabled,
            node_debug_enabled=node_debug_enabled,
            timing_measurement_enabled=bool(performance.get("measure_timing", True)),
            publish_timing_topic=bool(performance.get("publish_timing", True)),
            timing_topic=str(performance.get("timing_topic", "/rsg/preprocessor/timing")),
            write_timing_excel=bool(performance.get("write_timing_excel", True)),
            timing_excel_path=timing_excel_path,
            timing_sheet_name=timing_sheet_name,
            timing_excel_autosave_every=int(performance.get("timing_excel_autosave_every", 0)),
        )


class TransformMath:
    """Rigid-body transform utilities used by the preprocessing node."""

    @staticmethod
    def quaternion_to_matrix(qx: float, qy: float, qz: float, qw: float) -> np.ndarray:
        """Convert a quaternion to a 3x3 rotation matrix."""
        norm = math.sqrt(qx * qx + qy * qy + qz * qz + qw * qw)
        if norm == 0.0:
            return np.eye(3, dtype=np.float64)
        qx, qy, qz, qw = qx / norm, qy / norm, qz / norm, qw / norm

        return np.array([
            [1.0 - 2.0 * (qy * qy + qz * qz), 2.0 * (qx * qy - qz * qw), 2.0 * (qx * qz + qy * qw)],
            [2.0 * (qx * qy + qz * qw), 1.0 - 2.0 * (qx * qx + qz * qz), 2.0 * (qy * qz - qx * qw)],
            [2.0 * (qx * qz - qy * qw), 2.0 * (qy * qz + qx * qw), 1.0 - 2.0 * (qx * qx + qy * qy)],
        ], dtype=np.float64)

    @staticmethod
    def matrix_to_quaternion(rotation: np.ndarray) -> Tuple[float, float, float, float]:
        """Convert a 3x3 rotation matrix to a quaternion in x, y, z, w order."""
        m = rotation
        trace = float(np.trace(m))
        if trace > 0.0:
            s = math.sqrt(trace + 1.0) * 2.0
            qw = 0.25 * s
            qx = (m[2, 1] - m[1, 2]) / s
            qy = (m[0, 2] - m[2, 0]) / s
            qz = (m[1, 0] - m[0, 1]) / s
        elif m[0, 0] > m[1, 1] and m[0, 0] > m[2, 2]:
            s = math.sqrt(1.0 + m[0, 0] - m[1, 1] - m[2, 2]) * 2.0
            qw = (m[2, 1] - m[1, 2]) / s
            qx = 0.25 * s
            qy = (m[0, 1] + m[1, 0]) / s
            qz = (m[0, 2] + m[2, 0]) / s
        elif m[1, 1] > m[2, 2]:
            s = math.sqrt(1.0 + m[1, 1] - m[0, 0] - m[2, 2]) * 2.0
            qw = (m[0, 2] - m[2, 0]) / s
            qx = (m[0, 1] + m[1, 0]) / s
            qy = 0.25 * s
            qz = (m[1, 2] + m[2, 1]) / s
        else:
            s = math.sqrt(1.0 + m[2, 2] - m[0, 0] - m[1, 1]) * 2.0
            qw = (m[1, 0] - m[0, 1]) / s
            qx = (m[0, 2] + m[2, 0]) / s
            qy = (m[1, 2] + m[2, 1]) / s
            qz = 0.25 * s

        norm = math.sqrt(qx * qx + qy * qy + qz * qz + qw * qw)
        return qx / norm, qy / norm, qz / norm, qw / norm

    @staticmethod
    def rpy_to_matrix(roll: float, pitch: float, yaw: float) -> np.ndarray:
        """Convert roll, pitch, yaw angles to a rotation matrix.

        The convention is ``R = Rz(yaw) @ Ry(pitch) @ Rx(roll)``.
        """
        cr, sr = math.cos(roll), math.sin(roll)
        cp, sp = math.cos(pitch), math.sin(pitch)
        cy, sy = math.cos(yaw), math.sin(yaw)

        rx = np.array([[1, 0, 0], [0, cr, -sr], [0, sr, cr]], dtype=np.float64)
        ry = np.array([[cp, 0, sp], [0, 1, 0], [-sp, 0, cp]], dtype=np.float64)
        rz = np.array([[cy, -sy, 0], [sy, cy, 0], [0, 0, 1]], dtype=np.float64)
        return rz @ ry @ rx

    @staticmethod
    def make_transform(rotation: np.ndarray, translation: np.ndarray) -> np.ndarray:
        """Create a 4x4 homogeneous transform from rotation and translation."""
        transform = np.eye(4, dtype=np.float64)
        transform[:3, :3] = rotation
        transform[:3, 3] = translation.reshape(3)
        return transform

    @staticmethod
    def odom_to_transform(odom_msg: Odometry) -> np.ndarray:
        """Convert a ROS Odometry message into a homogeneous transform."""
        pose = odom_msg.pose.pose
        translation = np.array([pose.position.x, pose.position.y, pose.position.z], dtype=np.float64)
        rotation = TransformMath.quaternion_to_matrix(
            pose.orientation.x,
            pose.orientation.y,
            pose.orientation.z,
            pose.orientation.w,
        )
        return TransformMath.make_transform(rotation, translation)

    @staticmethod
    def slerp(q0: np.ndarray, q1: np.ndarray, alpha: float) -> np.ndarray:
        """Spherically interpolate two quaternions in x, y, z, w order."""
        q0 = q0 / np.linalg.norm(q0)
        q1 = q1 / np.linalg.norm(q1)
        dot = float(np.dot(q0, q1))

        if dot < 0.0:
            q1 = -q1
            dot = -dot

        if dot > 0.9995:
            result = q0 + alpha * (q1 - q0)
            return result / np.linalg.norm(result)

        theta_0 = math.acos(dot)
        theta = theta_0 * alpha
        sin_theta = math.sin(theta)
        sin_theta_0 = math.sin(theta_0)

        s0 = math.cos(theta) - dot * sin_theta / sin_theta_0
        s1 = sin_theta / sin_theta_0
        return (s0 * q0) + (s1 * q1)


class OdomBuffer:
    """Time-ordered odometry buffer with nearest and interpolated lookup."""

    def __init__(self, max_size: int, tolerance_sec: float, use_interpolation: bool) -> None:
        """Initialize the odometry buffer.

        Args:
            max_size: Maximum number of odometry messages to keep.
            tolerance_sec: Maximum allowed odometry-to-RGB timestamp difference.
            use_interpolation: Whether to interpolate between odometry samples.
        """
        self.max_size = max_size
        self.tolerance_sec = tolerance_sec
        self.use_interpolation = use_interpolation
        self._messages: List[Odometry] = []

    def add(self, msg: Odometry) -> None:
        """Add an odometry message to the buffer."""
        self._messages.append(msg)
        self._messages.sort(key=lambda m: stamp_to_float(m.header.stamp))
        if len(self._messages) > self.max_size:
            self._messages = self._messages[-self.max_size:]

    def lookup(self, target_time: float) -> Tuple[Optional[np.ndarray], Optional[float], str]:
        """Find or interpolate odometry transform for a target time.

        Args:
            target_time: RGB timestamp in seconds.

        Returns:
            Tuple containing the homogeneous transform, odometry timestamp delta,
            and a status string.
        """
        if not self._messages:
            return None, None, "odom_buffer_empty"

        if self.use_interpolation:
            interpolated = self._lookup_interpolated(target_time)
            if interpolated[0] is not None:
                return interpolated

        return self._lookup_nearest(target_time)

    def _lookup_nearest(self, target_time: float) -> Tuple[Optional[np.ndarray], Optional[float], str]:
        nearest = min(self._messages, key=lambda m: abs(stamp_to_float(m.header.stamp) - target_time))
        delta = abs(stamp_to_float(nearest.header.stamp) - target_time)
        if delta > self.tolerance_sec:
            return None, delta, "nearest_odom_outside_tolerance"
        return TransformMath.odom_to_transform(nearest), delta, "nearest_odom"

    def _lookup_interpolated(self, target_time: float) -> Tuple[Optional[np.ndarray], Optional[float], str]:
        before: Optional[Odometry] = None
        after: Optional[Odometry] = None

        for msg in self._messages:
            msg_time = stamp_to_float(msg.header.stamp)
            if msg_time <= target_time:
                before = msg
            elif msg_time > target_time:
                after = msg
                break

        if before is None or after is None:
            return None, None, "interpolation_bounds_missing"

        before_time = stamp_to_float(before.header.stamp)
        after_time = stamp_to_float(after.header.stamp)
        nearest_delta = min(abs(target_time - before_time), abs(after_time - target_time))

        if nearest_delta > self.tolerance_sec:
            return None, nearest_delta, "interpolated_odom_outside_tolerance"

        if after_time <= before_time:
            return TransformMath.odom_to_transform(before), 0.0, "duplicate_odom_time"

        alpha = (target_time - before_time) / (after_time - before_time)

        p0 = before.pose.pose.position
        p1 = after.pose.pose.position
        translation = np.array([
            p0.x + alpha * (p1.x - p0.x),
            p0.y + alpha * (p1.y - p0.y),
            p0.z + alpha * (p1.z - p0.z),
        ], dtype=np.float64)

        o0 = before.pose.pose.orientation
        o1 = after.pose.pose.orientation
        q0 = np.array([o0.x, o0.y, o0.z, o0.w], dtype=np.float64)
        q1 = np.array([o1.x, o1.y, o1.z, o1.w], dtype=np.float64)
        q = TransformMath.slerp(q0, q1, alpha)
        rotation = TransformMath.quaternion_to_matrix(q[0], q[1], q[2], q[3])

        return TransformMath.make_transform(rotation, translation), nearest_delta, "interpolated_odom"


class ImuBuffer:
    """Time-ordered IMU buffer with nearest-sample lookup."""

    def __init__(self, max_size: int, tolerance_sec: float) -> None:
        """Initialize the IMU buffer.

        Args:
            max_size: Maximum number of IMU messages to keep.
            tolerance_sec: Maximum allowed IMU-to-RGB timestamp difference.
        """
        self.max_size = max_size
        self.tolerance_sec = tolerance_sec
        self._messages: List[Imu] = []

    def add(self, msg: Imu) -> None:
        """Add an IMU message to the buffer."""
        self._messages.append(msg)
        self._messages.sort(key=lambda m: stamp_to_float(m.header.stamp))
        if len(self._messages) > self.max_size:
            self._messages = self._messages[-self.max_size:]

    def lookup(self, target_time: float) -> Tuple[Optional[Imu], Optional[float], str]:
        """Find the nearest IMU sample for a target RGB timestamp.

        Args:
            target_time: RGB timestamp in seconds.

        Returns:
            Tuple containing the nearest IMU message, timestamp delta, and a
            status string. If no valid IMU sample is available within the
            configured tolerance, the message is ``None``.
        """
        if not self._messages:
            return None, None, "imu_buffer_empty"

        nearest = min(self._messages, key=lambda m: abs(stamp_to_float(m.header.stamp) - target_time))
        delta = abs(stamp_to_float(nearest.header.stamp) - target_time)
        if delta > self.tolerance_sec:
            return None, delta, "nearest_imu_outside_tolerance"
        return nearest, delta, "nearest_imu"


class ImageConverter:
    """Convert ROS image messages into normalized ROS image messages and arrays."""

    def __init__(self, bridge: CvBridge, config: PreprocessorConfig) -> None:
        """Initialize the image converter."""
        self.bridge = bridge
        self.config = config

    def convert_rgb(self, rgb_msg: Image) -> Tuple[np.ndarray, Image]:
        """Convert an input image message to RGB uint8 format.

        Args:
            rgb_msg: Input color image message.

        Returns:
            Tuple of RGB NumPy array and converted ROS image message.
        """
        rgb_array = self.bridge.imgmsg_to_cv2(rgb_msg, desired_encoding="rgb8")
        if rgb_array.dtype != np.uint8:
            rgb_array = np.clip(rgb_array, 0, 255).astype(np.uint8)

        out_msg = self.bridge.cv2_to_imgmsg(rgb_array, encoding="rgb8")
        out_msg.header = copy.deepcopy(rgb_msg.header)
        out_msg.header.frame_id = self.config.camera_frame
        return rgb_array, out_msg

    def convert_depth(self, depth_msg: Image) -> Tuple[np.ndarray, np.ndarray, Image, float]:
        """Convert aligned depth to metres and create a validity mask.

        Args:
            depth_msg: Input aligned depth image message.

        Returns:
            Tuple containing metric depth image, validity mask, converted depth
            message, and invalid depth ratio.

        Raises:
            ValueError: If the input depth encoding is unsupported.
        """
        depth_raw = self.bridge.imgmsg_to_cv2(depth_msg, desired_encoding="passthrough")

        if depth_msg.encoding == "16UC1":
            depth_m = depth_raw.astype(np.float32) * self.config.depth_scale_to_meter
        elif depth_msg.encoding == "32FC1":
            depth_m = depth_raw.astype(np.float32)
        else:
            raise ValueError(f"Unsupported depth encoding: {depth_msg.encoding}")

        valid_mask = (
            np.isfinite(depth_m)
            & (depth_m >= self.config.min_depth_m)
            & (depth_m <= self.config.max_depth_m)
        )
        invalid_ratio = 1.0 - float(np.count_nonzero(valid_mask)) / float(valid_mask.size)

        out_msg = self.bridge.cv2_to_imgmsg(depth_m, encoding="32FC1")
        out_msg.header = copy.deepcopy(depth_msg.header)
        out_msg.header.frame_id = self.config.camera_frame
        return depth_m, valid_mask, out_msg, invalid_ratio


class FrameValidator:
    """Validation rules for synchronized RGB-D-pose frames."""

    def __init__(self, config: PreprocessorConfig) -> None:
        """Initialize the frame validator."""
        self.config = config

    def validate_resolution(self, rgb_msg: Image, depth_msg: Image) -> Optional[str]:
        """Validate that RGB and aligned depth resolutions match."""
        if not self.config.reject_resolution_mismatch:
            return None
        if rgb_msg.width != depth_msg.width or rgb_msg.height != depth_msg.height:
            return "rgb_depth_resolution_mismatch"
        return None

    def validate_depth_ratio(self, invalid_depth_ratio: float) -> Optional[str]:
        """Validate that the invalid-depth ratio is below the configured limit."""
        if invalid_depth_ratio > self.config.max_invalid_depth_ratio:
            return "too_many_invalid_depth_pixels"
        return None


class TimingExcelRecorder:
    """Collect preprocessing timing/debug events and export an Excel report.

    The recorder is inactive when debug mode is disabled. It stores one row for
    each accepted or rejected RGB-D synchronization event. Published frames are
    marked as ``published`` and rejected frames are marked as ``rejected`` with
    the rejection reason highlighted in the generated workbook.
    """

    def __init__(self, enabled: bool, output_path: str, autosave_every: int, logger: Any, sheet_name: str = "timing_debug") -> None:
        """Initialize the timing recorder.

        Args:
            enabled: Whether timing data should be recorded.
            output_path: Path of the Excel workbook to write.
            autosave_every: Save after this many samples. ``0`` disables autosave.
            logger: ROS logger used for status messages.
            sheet_name: Name of the worksheet created inside this node-specific workbook.
        """
        self.enabled = enabled
        self.output_path = Path(output_path).expanduser()
        self.autosave_every = max(0, int(autosave_every))
        self.logger = logger
        self.samples: List[Dict[str, Any]] = []
        self._last_saved_count = 0
        self.sheet_name = self._sanitize_sheet_name(sheet_name)
        if self.enabled:
            self.logger.info(f"Timing Excel recorder enabled. Output: {self.output_path}")
        else:
            self.logger.info("Timing Excel recorder disabled.")

    @staticmethod
    def _sanitize_sheet_name(name: str) -> str:
        """Return an Excel-safe worksheet name limited to 31 characters."""
        invalid_chars = "[]:*?/\\"
        cleaned = "".join("_" if char in invalid_chars else char for char in name)
        return (cleaned.strip() or "timing_debug")[:31]

    def add_sample(
        self,
        sequence: int,
        frame_id: str,
        rgb_time: float,
        processing_delay_ms: float,
        rgb_depth_dt_sec: Optional[float],
        rgb_odom_dt_sec: Optional[float],
        odom_status: str,
        status: str = "published",
        reason: str = "ok",
        invalid_depth_ratio: Optional[float] = None,
        imu_status: Optional[str] = None,
        rgb_imu_dt_sec: Optional[float] = None,
    ) -> None:
        """Add one preprocessing timing/debug event.

        Args:
            sequence: Sequential frame number generated by the preprocessor.
            frame_id: Generated Risk Scene Graph frame identifier.
            rgb_time: RGB frame timestamp in seconds.
            processing_delay_ms: Wall-clock preprocessing delay in milliseconds.
            rgb_depth_dt_sec: Timestamp difference between RGB and depth.
            rgb_odom_dt_sec: Timestamp difference between RGB and odometry.
            odom_status: Odometry association method or rejection status.
            status: ``published`` for accepted frames or ``rejected`` for rejected frames.
            reason: Rejection reason or ``ok`` for published frames.
            invalid_depth_ratio: Fraction of invalid depth pixels, if available.
            imu_status: IMU association status, if available.
            rgb_imu_dt_sec: Timestamp difference between RGB and IMU, if available.
        """
        if not self.enabled:
            return

        self.samples.append({
            "sequence": int(sequence),
            "frame_id": frame_id,
            "status": status,
            "reason": reason,
            "rgb_time_sec": float(rgb_time),
            "processing_delay_ms": float(processing_delay_ms),
            "rgb_depth_dt_sec": None if rgb_depth_dt_sec is None else float(rgb_depth_dt_sec),
            "rgb_odom_dt_sec": None if rgb_odom_dt_sec is None else float(rgb_odom_dt_sec),
            "odom_status": odom_status,
            "invalid_depth_ratio": None if invalid_depth_ratio is None else float(invalid_depth_ratio),
            "imu_status": imu_status,
            "rgb_imu_dt_sec": None if rgb_imu_dt_sec is None else float(rgb_imu_dt_sec),
        })

        # Create the Excel file as soon as the first timing/debug event is
        # available so users can immediately verify that debug logging is
        # active. After that, update it every autosave_every events when
        # autosave is enabled.
        if len(self.samples) == 1:
            self.save()
        elif self.autosave_every > 0 and len(self.samples) % self.autosave_every == 0:
            self.save()

    def save(self) -> None:
        """Write the timing/debug table and delay chart to an Excel workbook."""
        if not self.enabled or not self.samples:
            return
        if self._last_saved_count == len(self.samples) and self.output_path.exists():
            return

        try:
            from openpyxl import Workbook
            from openpyxl.chart import LineChart, Reference
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            self.logger.error(
                "Could not write timing Excel file because openpyxl is missing. "
                "Install it with: sudo apt install python3-openpyxl or python3 -m pip install openpyxl"
            )
            self.logger.error(str(exc))
            return

        self.output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.sheet_name

        headers = [
            "Frame sequence",
            "Frame ID",
            "Status",
            "Reason",
            "RGB timestamp [s]",
            "Preprocessing delay [ms]",
            "RGB-depth dt [s]",
            "RGB-odom dt [s]",
            "Odom status",
            "Invalid depth ratio",
            "IMU status",
            "RGB-IMU dt [s]",
        ]
        sheet.append(headers)

        for sample in self.samples:
            sheet.append([
                sample["sequence"],
                sample["frame_id"],
                sample["status"],
                sample["reason"],
                sample["rgb_time_sec"],
                sample["processing_delay_ms"],
                sample["rgb_depth_dt_sec"],
                sample["rgb_odom_dt_sec"],
                sample["odom_status"],
                sample["invalid_depth_ratio"],
                sample["imu_status"],
                sample["rgb_imu_dt_sec"],
            ])

        header_fill = PatternFill("solid", fgColor="D9EAF7")
        rejected_fill = PatternFill("solid", fgColor="F4CCCC")
        published_fill = PatternFill("solid", fgColor="D9EAD3")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row in sheet.iter_rows(min_row=2, max_row=len(self.samples) + 1):
            status_value = str(row[2].value).lower() if row[2].value is not None else ""
            if status_value == "rejected":
                for cell in row:
                    cell.fill = rejected_fill
                row[2].font = Font(bold=True, color="9C0006")
                row[3].font = Font(bold=True, color="9C0006")
            elif status_value == "published":
                row[2].fill = published_fill

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = {
            "A": 16,
            "B": 42,
            "C": 14,
            "D": 34,
            "E": 18,
            "F": 24,
            "G": 18,
            "H": 18,
            "I": 26,
            "J": 20,
            "K": 22,
            "L": 18,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        for row in sheet.iter_rows(min_row=2, min_col=5, max_col=12):
            for cell in row:
                cell.number_format = "0.000000"
        for cell in sheet["F"][1:]:
            cell.number_format = "0.000"
        for cell in sheet["J"][1:]:
            cell.number_format = "0.000"

        if len(self.samples) >= 2:
            chart = LineChart()
            chart.title = "Preprocessing delay per RGB frame"
            chart.style = 13
            chart.y_axis.title = "Delay [ms]"
            chart.x_axis.title = "Incoming RGB frame sequence"

            data = Reference(sheet, min_col=6, min_row=1, max_row=len(self.samples) + 1)
            categories = Reference(sheet, min_col=1, min_row=2, max_row=len(self.samples) + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            chart.height = 12
            chart.width = 24
            sheet.add_chart(chart, "N2")

        workbook.save(self.output_path)
        self._last_saved_count = len(self.samples)
        self.logger.info(f"Wrote preprocessing timing Excel report: {self.output_path}")


class RSGPreProcessor(Node):
    """ROS 2 node that prepares synchronized RGB-D-pose frames for RSG workers.

    The node keeps the downstream interface close to the original framework by
    publishing all prepared data in a single ``RsgFrame`` message. The main
    worker node can convert this message into the existing dictionary format and
    maintain its own processing queue.
    """

    def __init__(self) -> None:
        """Initialize subscriptions, publishers, configuration, and buffers."""
        super().__init__("RSG_pre_processor")

        self.declare_parameter("config_file", "")
        config_file = self.get_parameter("config_file").get_parameter_value().string_value
        if not config_file:
            raise ValueError("Parameter 'config_file' must point to rsg_pipeline.yaml")

        self.config = PreprocessorConfig.from_yaml(config_file)
        self.set_parameters([rclpy.parameter.Parameter("use_sim_time", rclpy.Parameter.Type.BOOL, self.config.use_sim_time)])

        self.bridge = CvBridge()
        self.converter = ImageConverter(self.bridge, self.config)
        self.validator = FrameValidator(self.config)
        self.odom_buffer = OdomBuffer(
            max_size=self.config.odom_buffer_size,
            tolerance_sec=self.config.odom_tolerance_sec,
            use_interpolation=self.config.use_odom_interpolation,
        )
        self.imu_buffer: Optional[ImuBuffer] = None
        if self.config.use_camera_imu:
            self.imu_buffer = ImuBuffer(
                max_size=self.config.imu_buffer_size,
                tolerance_sec=self.config.imu_tolerance_sec,
            )

        base_rotation = TransformMath.rpy_to_matrix(*self.config.base_to_camera_rpy)
        base_translation = np.array(self.config.base_to_camera_translation, dtype=np.float64)
        self.t_base_camera = TransformMath.make_transform(base_rotation, base_translation)

        sensor_qos = qos_profile_sensor_data
        sensor_qos.depth = self.config.sensor_qos_depth
        output_qos = QoSProfile(
            history=HistoryPolicy.KEEP_LAST,
            depth=self.config.output_qos_depth,
            reliability=ReliabilityPolicy.RELIABLE,
        )
        odom_qos = QoSProfile(
            history=HistoryPolicy.KEEP_LAST,
            depth=self.config.odom_qos_depth,
            reliability=ReliabilityPolicy.RELIABLE,
        )

        self.rgb_sub = message_filters.Subscriber(self, Image, self.config.rgb_topic, qos_profile=sensor_qos)
        self.depth_sub = message_filters.Subscriber(self, Image, self.config.depth_topic, qos_profile=sensor_qos)
        self.info_sub = message_filters.Subscriber(self, CameraInfo, self.config.camera_info_topic, qos_profile=sensor_qos)
        self.sync = message_filters.ApproximateTimeSynchronizer(
            [self.rgb_sub, self.depth_sub, self.info_sub],
            queue_size=30,
            slop=self.config.rgb_depth_slop_sec,
        )
        self.sync.registerCallback(self.synced_callback)

        self.odom_sub = self.create_subscription(Odometry, self.config.odom_topic, self.odom_callback, odom_qos)
        self.imu_sub = None
        if self.config.use_camera_imu:
            self.imu_sub = self.create_subscription(Imu, self.config.imu_topic, self.imu_callback, sensor_qos)
        self.frame_pub = self.create_publisher(RsgFrame, self.config.frame_topic, output_qos)
        self.status_pub = self.create_publisher(String, self.config.status_topic, output_qos)
        self.timing_pub = None
        if self.config.timing_enabled and self.config.publish_timing_topic:
            self.timing_pub = self.create_publisher(Float64MultiArray, self.config.timing_topic, output_qos)

        self.timing_recorder = TimingExcelRecorder(
            enabled=self.config.timing_enabled and self.config.write_timing_excel,
            output_path=self.config.timing_excel_path,
            autosave_every=self.config.timing_excel_autosave_every,
            logger=self.get_logger(),
            sheet_name=self.config.timing_sheet_name,
        )

        self.sequence = 0
        self.get_logger().info("RSG_pre_processor started.")
        self.get_logger().info(f"RGB topic: {self.config.rgb_topic}")
        self.get_logger().info(f"Depth topic: {self.config.depth_topic}")
        self.get_logger().info(f"CameraInfo topic: {self.config.camera_info_topic}")
        self.get_logger().info(f"Odom topic: {self.config.odom_topic}")
        if self.config.use_camera_imu:
            self.get_logger().info(f"Camera IMU topic: {self.config.imu_topic}")
            self.get_logger().info(
                f"Camera IMU enabled: require_for_frame={self.config.require_imu_for_frame}, "
                f"tolerance={self.config.imu_tolerance_sec}s"
            )
        else:
            self.get_logger().info("Camera IMU support disabled in configuration.")
        self.get_logger().info(f"Output frame topic: {self.config.frame_topic}")
        self.get_logger().info(
            f"Debug enabled: {self.config.debug_enabled} "
            f"(global={self.config.global_debug_enabled}, "
            f"node={self.config.node_debug_enabled})"
        )
        self.get_logger().info(
            f"Timing enabled: {self.config.timing_enabled} "
            f"(measure={self.config.timing_measurement_enabled}, "
            f"write_excel={self.config.write_timing_excel}, "
            f"publish_topic={self.config.publish_timing_topic})"
        )
        if self.config.timing_enabled:
            self.get_logger().info(f"Timing measurement enabled. Excel debug file: {self.config.timing_excel_path}")

    def odom_callback(self, msg: Odometry) -> None:
        """Store incoming odometry messages for timestamp association."""
        self.odom_buffer.add(msg)

    def imu_callback(self, msg: Imu) -> None:
        """Store incoming camera IMU messages for timestamp association."""
        if self.imu_buffer is not None:
            self.imu_buffer.add(msg)

    def synced_callback(self, rgb_msg: Image, depth_msg: Image, info_msg: CameraInfo) -> None:
        """Prepare and publish one synchronized RGB-D-pose frame.

        Args:
            rgb_msg: Synchronized RGB image message.
            depth_msg: Synchronized aligned depth image message.
            info_msg: Synchronized camera calibration message.
        """
        processing_start = time.perf_counter() if self.config.timing_enabled else None

        self.sequence += 1
        rgb_time = stamp_to_float(rgb_msg.header.stamp)
        depth_time = stamp_to_float(depth_msg.header.stamp)
        info_time = stamp_to_float(info_msg.header.stamp)
        rgb_depth_dt = abs(rgb_time - depth_time)
        rgb_camera_info_dt = abs(rgb_time - info_time)

        frame_id = self.make_frame_id(rgb_msg)

        resolution_error = self.validator.validate_resolution(rgb_msg, depth_msg)
        if resolution_error:
            self.reject_frame(
                frame_id=frame_id,
                reason=resolution_error,
                rgb_time=rgb_time,
                processing_start=processing_start,
                extra={
                    "rgb_depth_dt_sec": rgb_depth_dt,
                    "rgb_camera_info_dt_sec": rgb_camera_info_dt,
                },
                rgb_depth_dt_sec=rgb_depth_dt,
            )
            return

        try:
            _rgb_array, rgb_out = self.converter.convert_rgb(rgb_msg)
            _depth_m, _valid_mask, depth_out, invalid_depth_ratio = self.converter.convert_depth(depth_msg)
        except Exception as exc:
            self.reject_frame(
                frame_id=frame_id,
                reason=str(exc),
                rgb_time=rgb_time,
                processing_start=processing_start,
                extra={
                    "rgb_depth_dt_sec": rgb_depth_dt,
                    "rgb_camera_info_dt_sec": rgb_camera_info_dt,
                },
                rgb_depth_dt_sec=rgb_depth_dt,
            )
            return

        depth_error = self.validator.validate_depth_ratio(invalid_depth_ratio)
        if depth_error:
            self.reject_frame(
                frame_id=frame_id,
                reason=depth_error,
                rgb_time=rgb_time,
                processing_start=processing_start,
                extra={
                    "invalid_depth_ratio": invalid_depth_ratio,
                    "rgb_depth_dt_sec": rgb_depth_dt,
                    "rgb_camera_info_dt_sec": rgb_camera_info_dt,
                },
                rgb_depth_dt_sec=rgb_depth_dt,
                invalid_depth_ratio=invalid_depth_ratio,
            )
            return

        t_odom_base, odom_delta, odom_status = self.odom_buffer.lookup(rgb_time)
        if t_odom_base is None:
            self.reject_frame(
                frame_id=frame_id,
                reason=odom_status,
                rgb_time=rgb_time,
                processing_start=processing_start,
                extra={
                    "odom_delta_sec": odom_delta,
                    "rgb_depth_dt_sec": rgb_depth_dt,
                    "rgb_camera_info_dt_sec": rgb_camera_info_dt,
                    "odom_status": odom_status,
                    "invalid_depth_ratio": invalid_depth_ratio,
                },
                rgb_depth_dt_sec=rgb_depth_dt,
                rgb_odom_dt_sec=odom_delta,
                odom_status=odom_status,
                invalid_depth_ratio=invalid_depth_ratio,
            )
            return

        imu_msg, imu_delta, imu_status = self.lookup_camera_imu(rgb_msg, rgb_time)
        has_camera_imu = imu_msg is not None
        if self.config.use_camera_imu and self.config.require_imu_for_frame and not has_camera_imu:
            self.reject_frame(
                frame_id=frame_id,
                reason=imu_status,
                rgb_time=rgb_time,
                processing_start=processing_start,
                extra={
                    "imu_delta_sec": imu_delta,
                    "rgb_imu_dt_sec": imu_delta,
                    "imu_status": imu_status,
                    "rgb_depth_dt_sec": rgb_depth_dt,
                    "rgb_camera_info_dt_sec": rgb_camera_info_dt,
                    "rgb_odom_dt_sec": odom_delta,
                    "odom_status": odom_status,
                    "invalid_depth_ratio": invalid_depth_ratio,
                },
                rgb_depth_dt_sec=rgb_depth_dt,
                rgb_odom_dt_sec=odom_delta,
                odom_status=odom_status,
                invalid_depth_ratio=invalid_depth_ratio,
                imu_status=imu_status,
                rgb_imu_dt_sec=imu_delta,
            )
            return

        t_odom_camera = t_odom_base @ self.t_base_camera
        tx = t_odom_camera[:3, 3].astype(np.float64)
        rot_m = t_odom_camera[:3, :3].astype(np.float64)

        pose_msg = self.make_pose_msg(rgb_msg, tx, rot_m)
        camera_info_out = copy.deepcopy(info_msg)
        camera_info_out.header.stamp = rgb_msg.header.stamp
        camera_info_out.header.frame_id = self.config.camera_frame

        rgb_out.header.stamp = rgb_msg.header.stamp
        rgb_out.header.frame_id = self.config.camera_frame
        depth_out.header.stamp = rgb_msg.header.stamp
        depth_out.header.frame_id = self.config.camera_frame

        msg = RsgFrame()
        msg.header.stamp = rgb_msg.header.stamp
        msg.header.frame_id = self.config.world_frame
        msg.rsg_frame_id = frame_id
        msg.source = self.config.source
        msg.sequence = self.sequence
        msg.rgb = rgb_out
        msg.depth_m = depth_out
        msg.camera_info = camera_info_out
        msg.has_camera_imu = bool(has_camera_imu)
        msg.camera_imu = imu_msg if imu_msg is not None else self.make_empty_imu_msg(rgb_msg)
        msg.rgb_imu_dt_sec = float(imu_delta if imu_delta is not None else -1.0)
        msg.camera_pose = pose_msg
        msg.tx = [float(v) for v in tx]
        msg.rot_m = [float(v) for v in rot_m.reshape(9)]
        msg.invalid_depth_ratio = float(invalid_depth_ratio)
        msg.rgb_depth_dt_sec = float(rgb_depth_dt)
        msg.rgb_camera_info_dt_sec = float(rgb_camera_info_dt)
        msg.rgb_odom_dt_sec = float(odom_delta if odom_delta is not None else -1.0)
        msg.metadata_json = json.dumps({
            "odom_status": odom_status,
            "imu_enabled": self.config.use_camera_imu,
            "imu_required": self.config.require_imu_for_frame,
            "imu_status": imu_status,
            "has_camera_imu": has_camera_imu,
            "imu_topic": self.config.imu_topic,
            "rgb_encoding_in": rgb_msg.encoding,
            "depth_encoding_in": depth_msg.encoding,
            "rgb_frame_id_in": rgb_msg.header.frame_id,
            "depth_frame_id_in": depth_msg.header.frame_id,
            "camera_frame_out": self.config.camera_frame,
            "world_frame": self.config.world_frame,
            "base_frame": self.config.base_frame,
            "session_date": self.config.session_date,
        })

        self.frame_pub.publish(msg)

        processing_delay_ms: Optional[float] = None
        if self.config.timing_enabled and processing_start is not None:
            processing_delay_ms = (time.perf_counter() - processing_start) * 1000.0
            self.publish_timing(
                frame_id=frame_id,
                rgb_time=rgb_time,
                processing_delay_ms=processing_delay_ms,
                rgb_depth_dt_sec=rgb_depth_dt,
                rgb_odom_dt_sec=float(odom_delta if odom_delta is not None else -1.0),
                odom_status=odom_status,
                invalid_depth_ratio=invalid_depth_ratio,
                imu_status=imu_status,
                rgb_imu_dt_sec=imu_delta,
            )

        status_extra: Dict[str, Any] = {
            "invalid_depth_ratio": invalid_depth_ratio,
            "rgb_depth_dt_sec": rgb_depth_dt,
            "rgb_camera_info_dt_sec": rgb_camera_info_dt,
            "rgb_odom_dt_sec": odom_delta,
            "odom_status": odom_status,
            "imu_enabled": self.config.use_camera_imu,
            "has_camera_imu": has_camera_imu,
            "rgb_imu_dt_sec": imu_delta,
            "imu_status": imu_status,
        }
        if processing_delay_ms is not None:
            status_extra["processing_delay_ms"] = processing_delay_ms

        self.publish_status("published", frame_id, "ok", rgb_time, status_extra)

    def reject_frame(
        self,
        frame_id: str,
        reason: str,
        rgb_time: float,
        processing_start: Optional[float],
        extra: Optional[Dict[str, Any]] = None,
        rgb_depth_dt_sec: Optional[float] = None,
        rgb_odom_dt_sec: Optional[float] = None,
        odom_status: str = "",
        invalid_depth_ratio: Optional[float] = None,
        imu_status: Optional[str] = None,
        rgb_imu_dt_sec: Optional[float] = None,
    ) -> None:
        """Record and publish a rejected-frame debug entry.

        Rejected frames are not published on ``/rsg/preprocessed/frame``. When
        debug timing is enabled, this method still writes a special ``rejected``
        row into the Excel debug report so rejection events are visible during
        performance analysis.
        """
        payload_extra: Dict[str, Any] = dict(extra or {})
        processing_delay_ms: Optional[float] = None

        if self.config.timing_enabled and processing_start is not None:
            processing_delay_ms = (time.perf_counter() - processing_start) * 1000.0
            payload_extra["processing_delay_ms"] = processing_delay_ms

            if rgb_depth_dt_sec is None:
                rgb_depth_dt_sec = payload_extra.get("rgb_depth_dt_sec")
            if rgb_odom_dt_sec is None:
                rgb_odom_dt_sec = payload_extra.get("rgb_odom_dt_sec", payload_extra.get("odom_delta_sec"))
            if not odom_status:
                odom_status = str(payload_extra.get("odom_status", reason))
            if invalid_depth_ratio is None:
                invalid_depth_ratio = payload_extra.get("invalid_depth_ratio")
            if imu_status is None:
                imu_status = payload_extra.get("imu_status")
            if rgb_imu_dt_sec is None:
                rgb_imu_dt_sec = payload_extra.get("rgb_imu_dt_sec", payload_extra.get("imu_delta_sec"))

            self.record_timing_event(
                frame_id=frame_id,
                rgb_time=rgb_time,
                processing_delay_ms=processing_delay_ms,
                rgb_depth_dt_sec=rgb_depth_dt_sec,
                rgb_odom_dt_sec=rgb_odom_dt_sec,
                odom_status=odom_status,
                status="rejected",
                reason=reason,
                invalid_depth_ratio=invalid_depth_ratio,
                imu_status=imu_status,
                rgb_imu_dt_sec=rgb_imu_dt_sec,
                publish_timing_topic=False,
            )

        self.publish_status("rejected", frame_id, reason, rgb_time, payload_extra)

    def lookup_camera_imu(self, rgb_msg: Image, rgb_time: float) -> Tuple[Optional[Imu], Optional[float], str]:
        """Look up the camera IMU sample associated with the RGB timestamp.

        The IMU stream is optional. If IMU support is disabled, this method
        returns a ``None`` message and the status ``imu_disabled``. If IMU
        support is enabled but ``require_imu_for_frame`` is false, frames are
        still published when no matching IMU sample is available.
        """
        if not self.config.use_camera_imu:
            return None, None, "imu_disabled"
        if self.imu_buffer is None:
            return None, None, "imu_buffer_not_initialized"
        return self.imu_buffer.lookup(rgb_time)

    def make_empty_imu_msg(self, rgb_msg: Image) -> Imu:
        """Create a timestamped empty IMU message when IMU is unavailable.

        The accompanying ``has_camera_imu`` field in ``RsgFrame`` indicates
        whether this placeholder should be used by downstream modules.
        """
        imu_msg = Imu()
        imu_msg.header.stamp = rgb_msg.header.stamp
        imu_msg.header.frame_id = self.config.camera_frame
        return imu_msg

    def publish_timing(
        self,
        frame_id: str,
        rgb_time: float,
        processing_delay_ms: float,
        rgb_depth_dt_sec: float,
        rgb_odom_dt_sec: float,
        odom_status: str,
        invalid_depth_ratio: Optional[float] = None,
        imu_status: Optional[str] = None,
        rgb_imu_dt_sec: Optional[float] = None,
    ) -> None:
        """Publish and record one successful preprocessing timing sample."""
        self.record_timing_event(
            frame_id=frame_id,
            rgb_time=rgb_time,
            processing_delay_ms=processing_delay_ms,
            rgb_depth_dt_sec=rgb_depth_dt_sec,
            rgb_odom_dt_sec=rgb_odom_dt_sec,
            odom_status=odom_status,
            status="published",
            reason="ok",
            invalid_depth_ratio=invalid_depth_ratio,
            imu_status=imu_status,
            rgb_imu_dt_sec=rgb_imu_dt_sec,
            publish_timing_topic=True,
        )

    def record_timing_event(
        self,
        frame_id: str,
        rgb_time: float,
        processing_delay_ms: float,
        rgb_depth_dt_sec: Optional[float],
        rgb_odom_dt_sec: Optional[float],
        odom_status: str,
        status: str,
        reason: str,
        invalid_depth_ratio: Optional[float] = None,
        imu_status: Optional[str] = None,
        rgb_imu_dt_sec: Optional[float] = None,
        publish_timing_topic: bool = True,
    ) -> None:
        """Publish timing data and/or record an Excel timing/debug row.

        Rejected frames are recorded in Excel with ``status='rejected'`` but
        are not published to the timing topic by default because no prepared
        frame was emitted for downstream processing.
        """
        if not self.config.timing_enabled:
            return

        if publish_timing_topic and self.timing_pub is not None:
            timing_msg = Float64MultiArray()
            timing_msg.data = [
                float(self.sequence),
                float(processing_delay_ms),
                float(rgb_time),
                float(rgb_depth_dt_sec if rgb_depth_dt_sec is not None else -1.0),
                float(rgb_odom_dt_sec if rgb_odom_dt_sec is not None else -1.0),
            ]
            self.timing_pub.publish(timing_msg)

        self.timing_recorder.add_sample(
            sequence=self.sequence,
            frame_id=frame_id,
            rgb_time=rgb_time,
            processing_delay_ms=processing_delay_ms,
            rgb_depth_dt_sec=rgb_depth_dt_sec,
            rgb_odom_dt_sec=rgb_odom_dt_sec,
            odom_status=odom_status,
            status=status,
            reason=reason,
            invalid_depth_ratio=invalid_depth_ratio,
            imu_status=imu_status,
            rgb_imu_dt_sec=rgb_imu_dt_sec,
        )

    def close(self) -> None:
        """Flush debug outputs before the node is destroyed."""
        self.timing_recorder.save()

    def make_frame_id(self, rgb_msg: Image) -> str:
        """Create a repeat-safe frame id from date, data timestamp, and sequence."""
        stamp = rgb_msg.header.stamp
        return (
            f"{self.config.frame_id_prefix}_"
            f"{self.config.session_date}_"
            f"{stamp.sec}_{stamp.nanosec:09d}_"
            f"{self.sequence:06d}"
        )

    def make_pose_msg(self, rgb_msg: Image, tx: np.ndarray, rot_m: np.ndarray) -> PoseStamped:
        """Create a camera pose message from translation and rotation matrix."""
        qx, qy, qz, qw = TransformMath.matrix_to_quaternion(rot_m)
        pose_msg = PoseStamped()
        pose_msg.header.stamp = rgb_msg.header.stamp
        pose_msg.header.frame_id = self.config.world_frame
        pose_msg.pose.position.x = float(tx[0])
        pose_msg.pose.position.y = float(tx[1])
        pose_msg.pose.position.z = float(tx[2])
        pose_msg.pose.orientation.x = float(qx)
        pose_msg.pose.orientation.y = float(qy)
        pose_msg.pose.orientation.z = float(qz)
        pose_msg.pose.orientation.w = float(qw)
        return pose_msg

    def publish_status(self, status: str, frame_id: str, reason: str, rgb_time: float, extra: Optional[Dict[str, Any]] = None) -> None:
        """Publish lightweight JSON status for debugging and bag inspection."""
        payload: Dict[str, Any] = {
            "status": status,
            "frame_id": frame_id,
            "reason": reason,
            "rgb_time": rgb_time,
            "sequence": self.sequence,
        }
        if extra:
            payload.update(extra)
        self.status_pub.publish(String(data=json.dumps(payload)))


def main(args: Optional[List[str]] = None) -> None:
    """Run the RSG_pre_processor node with graceful Ctrl+C handling."""
    rclpy.init(args=args)
    node: Optional[RSGPreProcessor] = None

    try:
        node = RSGPreProcessor()
        rclpy.spin(node)

    except KeyboardInterrupt:
        if node is not None:
            node.get_logger().info("Shutdown requested by user.")

    except rclpy.executors.ExternalShutdownException:
        # Raised when the launch system or another owner already shut down
        # the ROS context. This is normal during Ctrl+C from ros2 launch.
        pass

    finally:
        if node is not None:
            node.close()
            node.destroy_node()

        # During ros2 launch shutdown, rclpy may already have called shutdown.
        # Guarding with rclpy.ok() prevents: "rcl_shutdown already called".
        if rclpy.ok():
            rclpy.shutdown()


if __name__ == "__main__":
    main()

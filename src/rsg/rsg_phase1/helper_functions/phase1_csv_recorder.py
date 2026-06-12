"""Single-workbook Phase 1 debug recorder.

This module intentionally keeps the historical class name ``Phase1CsvDebugRecorder``
so the existing coordinator/classifier code does not need many import changes.
Internally it now writes one Excel workbook instead of several CSV files.

Workbook layout:
- HydraLatency
- ClassifierLatency
- FrameFIFO
- UnknownTracks
- VLMQueue
- VLMLatency

Rows are buffered in memory and saved on shutdown. Runtime autosave is disabled by default because repeatedly writing a growing XLSX workbook can add large timing spikes to the measured pipeline.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from threading import Lock, Thread
from typing import Any, Dict, Iterable, List, Optional, Tuple


class Phase1CsvDebugRecorder:
    """Write all simple Phase 1 debug tables into one XLSX workbook.

    The public method names are unchanged from the older CSV recorder:
    ``hydra_latency``, ``classifier_latency``, ``frame_fifo_event``,
    ``unknown_track_observation``, ``vlm_queue_event``, and ``vlm_latency``.
    This lets the node code stay simple while the output becomes one Excel file.
    """

    TABLES: Dict[str, Tuple[str, List[str]]] = {
        "phase1_cordinator_hydra_latency": (
            "HydraLatency",
            [
                "sequence",
                "frame_id",
                "status",
                "total_delay_ms",
                "sent_to_classifier_delay_ms",
                "coordinator_delay_ms",
                "classifier_delay_ms",
                "classifier_debug_record_delay_ms",
                "hydra_build_delay_ms",
                "hydra_publish_delay_ms",
                "unknown_publish_delay_ms",
                "evidence_record_delay_ms",
                "pipeline_wait_ms",
                "num_masks",
                "num_known",
                "num_unknown",
                "num_unknown_tracks",
                "num_vlm_queued",
            ],
        ),
        "phase1_classifier_phase_latency": (
            "ClassifierLatency",
            [
                "sequence",
                "frame_id",
                "status",
                "input_age_ms",
                "sam_delay_ms",
                "rap_delay_ms",
                "label_map_delay_ms",
                "metadata_delay_ms",
                "image_conversion_delay_ms",
                "result_message_build_delay_ms",
                "classifier_debug_record_delay_ms",
                "classifier_delay_ms",
                "num_masks",
                "num_known",
                "num_unknown",
                "num_new_tracks",
                "num_matched_tracks",
                "num_vlm_queued",
            ],
        ),
        "phase1_frame_fifo_queue": (
            "FrameFIFO",
            [
                "event_index",
                "event",
                "sequence",
                "frame_id",
                "queue_size",
                "queue_max_size",
                "queue_wait_ms",
                "reason",
            ],
        ),
        "phase1_unknown_tracks": (
            "UnknownTracks",
            [
                "sequence",
                "frame_id",
                "candidate_id",
                "unknown_track_id",
                "track_event",
                "track_seen_count",
                "vlm_status",
                "vlm_dispatch_status",
                "vlm_queue_size",
                "vlm_queue_max_size",
                "centroid_x",
                "centroid_y",
                "centroid_z",
                "bbox_volume_m3",
                "depth_valid_ratio",
                "mask_area_px",
                "best_frame_score",
            ],
        ),
        "phase1_vlm_queue": (
            "VLMQueue",
            [
                "event_index",
                "event",
                "sequence",
                "frame_id",
                "unknown_track_id",
                "candidate_id",
                "queue_size",
                "queue_max_size",
                "queue_wait_ms",
                "track_seen_count",
                "best_frame_score",
                "reason",
            ],
        ),
        "phase1_vlm_latency": (
            "VLMLatency",
            [
                "unknown_track_id",
                "candidate_id",
                "frame_id",
                "sequence",
                "status",
                "predicted_label",
                "confidence",
                "vlm_delay_ms",
                "total_age_ms",
                "track_seen_count",
                "best_frame_score",
                "backend",
                "model",
            ],
        ),
    }

    def __init__(self, enabled: bool, output_dir: str, node_name: str, logger: Any) -> None:
        self.enabled = enabled
        self.node_name = node_name
        self.logger = logger
        self.output_dir = Path(output_dir).expanduser().resolve()
        self.output_path = self._resolve_output_path(self.output_dir, node_name)
        self.autosave_every = 0

        self._lock = Lock()
        self._rows: Dict[str, List[Dict[str, Any]]] = {key: [] for key in self.TABLES.keys()}
        self._row_count = 0
        self._last_saved_count = 0
        self._save_thread: Optional[Thread] = None
        self._save_pending = False

        if self.enabled:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            self.logger.info(f"Phase 1 single Excel debug workbook enabled: {self.output_path}")

    @staticmethod
    def _resolve_output_path(path: Path, node_name: str) -> Path:
        """Treat a .xlsx path as a file; otherwise create one workbook in a folder."""
        if path.suffix.lower() == ".xlsx":
            return path
        return path / "phase1_debug.xlsx"

    def _append(self, table: str, row: Dict[str, Any]) -> None:
        if not self.enabled:
            return
        if table not in self.TABLES:
            return
        _, headers = self.TABLES[table]
        clean = {key: row.get(key, "") for key in headers}
        with self._lock:
            self._rows[table].append(clean)
            self._row_count += 1
            should_save = self.autosave_every > 0 and self._row_count % self.autosave_every == 0
        if should_save:
            self.request_save_async()

    def hydra_latency(self, **row: Any) -> None:
        self._append("phase1_cordinator_hydra_latency", row)

    def classifier_latency(self, **row: Any) -> None:
        self._append("phase1_classifier_phase_latency", row)

    def frame_fifo_event(self, **row: Any) -> None:
        self._append("phase1_frame_fifo_queue", row)

    def unknown_track_observation(self, **row: Any) -> None:
        self._append("phase1_unknown_tracks", row)

    def vlm_queue_event(self, **row: Any) -> None:
        self._append("phase1_vlm_queue", row)

    def vlm_latency(self, **row: Any) -> None:
        self._append("phase1_vlm_latency", row)

    def request_save_async(self) -> None:
        if not self.enabled:
            return
        with self._lock:
            if self._save_thread is not None and self._save_thread.is_alive():
                self._save_pending = True
                return
            self._save_pending = False
            self._save_thread = Thread(target=self._save_worker, daemon=True)
            self._save_thread.start()

    def _save_worker(self) -> None:
        while True:
            snapshot = self._snapshot_for_save()
            if snapshot is not None:
                rows, count = snapshot
                self._write_workbook(rows, count)
            with self._lock:
                if self._save_pending:
                    self._save_pending = False
                    continue
                return

    def _snapshot_for_save(self) -> Optional[Tuple[Dict[str, List[Dict[str, Any]]], int]]:
        with self._lock:
            count = self._row_count
            if count == 0:
                return None
            if count == self._last_saved_count and self.output_path.exists():
                return None
            return deepcopy(self._rows), count

    def close(self) -> None:
        if not self.enabled:
            return
        thread = self._save_thread
        if thread is not None and thread.is_alive():
            thread.join(timeout=2.0)
        snapshot = self._snapshot_for_save()
        if snapshot is not None:
            rows, count = snapshot
            self._write_workbook(rows, count)

    def _write_workbook(self, rows_by_table: Dict[str, List[Dict[str, Any]]], count: int) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.chart import LineChart, Reference
        except ImportError as exc:
            self.logger.error("openpyxl is missing; cannot write Phase 1 debug workbook.")
            self.logger.error(str(exc))
            return

        wb = Workbook()
        # Remove default sheet; we create all sheets explicitly.
        default = wb.active
        wb.remove(default)

        header_fill = PatternFill("solid", fgColor="D9EAF7")
        ok_fill = PatternFill("solid", fgColor="D9EAD3")
        warn_fill = PatternFill("solid", fgColor="FFF2CC")
        fail_fill = PatternFill("solid", fgColor="F4CCCC")

        for table_name, (sheet_name, headers) in self.TABLES.items():
            ws = wb.create_sheet(sheet_name)
            ws.append(headers)
            for row in rows_by_table.get(table_name, []):
                ws.append([row.get(header, "") for header in headers])

            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            # Highlight status/event cells for quick visual inspection.
            status_col = None
            event_col = None
            if "status" in headers:
                status_col = headers.index("status") + 1
            if "event" in headers:
                event_col = headers.index("event") + 1
            for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
                if status_col is not None:
                    value = str(row_cells[status_col - 1].value or "").lower()
                    if value in {"ok", "sent_to_hydra", "vlm_done", "completed"}:
                        row_cells[status_col - 1].fill = ok_fill
                    elif value in {"dropped", "failed", "error", "stale"}:
                        row_cells[status_col - 1].fill = fail_fill
                    elif value:
                        row_cells[status_col - 1].fill = warn_fill
                if event_col is not None:
                    value = str(row_cells[event_col - 1].value or "").lower()
                    if "dropped" in value or "failed" in value:
                        row_cells[event_col - 1].fill = fail_fill
                    elif value in {"completed", "dequeued", "dequeued_to_sam_rap", "dequeued_sent_to_classifier"}:
                        row_cells[event_col - 1].fill = ok_fill
                    elif value:
                        row_cells[event_col - 1].fill = warn_fill

            for col in ws.columns:
                letter = col[0].column_letter
                max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
                ws.column_dimensions[letter].width = min(max(12, max_len + 2), 44)

            # Add a simple chart on sheets where sequence/event_index and delay columns exist.
            if ws.max_row > 2 and sheet_name in {"HydraLatency", "ClassifierLatency", "FrameFIFO", "VLMQueue", "VLMLatency"}:
                self._add_basic_chart(ws, headers, LineChart, Reference)

        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.output_path)
        with self._lock:
            self._last_saved_count = count
        self.logger.info(f"Saved Phase 1 debug workbook with {count} rows: {self.output_path}")

    @staticmethod
    def _add_basic_chart(ws: Any, headers: List[str], LineChart: Any, Reference: Any) -> None:
        """Add a small chart to the right of the table for quick inspection."""
        try:
            x_key = "sequence" if "sequence" in headers else "event_index" if "event_index" in headers else None
            if x_key is None:
                return
            candidate_y_keys = [
                "total_delay_ms",
                "sent_to_classifier_delay_ms",
                "classifier_delay_ms",
                "coordinator_delay_ms",
                "pipeline_wait_ms",
                "sam_delay_ms",
                "rap_delay_ms",
                "label_map_delay_ms",
                "metadata_delay_ms",
                "queue_size",
                "queue_wait_ms",
                "vlm_delay_ms",
                "total_age_ms",
            ]
            y_keys = [key for key in candidate_y_keys if key in headers]
            if not y_keys:
                return
            chart = LineChart()
            chart.title = f"{ws.title} overview"
            chart.y_axis.title = "ms / count"
            chart.x_axis.title = x_key
            for key in y_keys[:6]:
                col_idx = headers.index(key) + 1
                data = Reference(ws, min_col=col_idx, min_row=1, max_row=ws.max_row)
                chart.add_data(data, titles_from_data=True)
            x_col = headers.index(x_key) + 1
            cats = Reference(ws, min_col=x_col, min_row=2, max_row=ws.max_row)
            chart.set_categories(cats)
            chart.height = 7
            chart.width = 16
            ws.add_chart(chart, f"{chr(65 + min(len(headers) + 1, 20))}2")
        except Exception:
            # Charts are only a convenience; never break debug writing because of a chart issue.
            return
